import csv
import shutil
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

import sync_hevo_training_tables as syncer


@pytest.fixture
def isolated_tables(tmp_path, monkeypatch):
    original = syncer.TABLES
    paths = [
        "Supplementary Tables.xlsx",
        "S46_architecture_and_hyperparameters_of_the_primary_hevo_models.csv",
        "S46_training_run_budgets.csv",
        "repository_tables/R19_hevo_architecture_and_training_configuration.tsv",
    ]
    for relative in paths:
        output = tmp_path / relative
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(original / relative, output)
    # Isolate blank-input behavior from the published, fully populated budgets.
    budget_path = tmp_path / "S46_training_run_budgets.csv"
    fields, records = syncer.read_table(budget_path)
    for record in records:
        if record["run"] in syncer.PUBLIC:
            for field in ["optimizer_step_budget", "sequence_budget", "nominal_nucleotide_positions"]:
                record[field] = ""
    syncer.write_table(budget_path, fields, records)
    workbook_path = tmp_path / "Supplementary Tables.xlsx"
    wb = load_workbook(workbook_path)
    for name in ["hEVO-base-o", "hEVO-large-o"]:
        row = public_row(wb["S30"], name)
        for col in [5, 6, 7]:
            wb["S30"].cell(row, col).value = None
    wb.save(workbook_path)
    monkeypatch.setattr(syncer, "TABLES", tmp_path)
    return tmp_path


def public_row(sheet, name="hEVO-base-o"):
    return next(row for row in range(1, sheet.max_row + 1)
                if sheet.cell(row, 1).value == name)


def test_blank_public_budgets_stay_blank_and_other_sheets_are_preserved(isolated_tables):
    workbook = isolated_tables / "Supplementary Tables.xlsx"
    with ZipFile(workbook) as archive:
        before = {n: archive.read(n) for n in archive.namelist()}
    assert syncer.sync(workbook) == 9
    ws = load_workbook(workbook)["S30"]
    for name in ["hEVO-base-o", "hEVO-large-o"]:
        row = public_row(ws, name)
        assert ws.cell(row, 5).value is None
        assert ws.cell(row, 6).value is None
        assert ws.cell(row, 7).value == f'=IF(F{row}="","",F{row}*1024)'
    with ZipFile(workbook) as archive:
        changed = [name for name in before if before[name] != archive.read(name)]
    assert all(name.endswith("sheet31.xml") for name in changed)


def test_entered_public_budgets_round_trip_to_csv_and_reader_table(isolated_tables):
    path = isolated_tables / "Supplementary Tables.xlsx"
    wb = load_workbook(path)
    row = public_row(wb["S30"])
    wb["S30"].cell(row, 5, 321)
    wb["S30"].cell(row, 6, 654321)
    wb.save(path)
    syncer.sync(path)
    _, entries = syncer.read_table(isolated_tables / "S46_training_run_budgets.csv")
    record = next(r for r in entries if r["run"] == "hEvo-open-base")
    assert record["optimizer_step_budget"] == "321"
    assert record["sequence_budget"] == "654321"
    assert record["nominal_nucleotide_positions"] == str(654321 * 1024)
    _, reader = syncer.read_table(isolated_tables / "repository_tables/R19_training_run_budgets.tsv", "\t")
    assert next(r for r in reader if r["run"] == "hEvo-open-base") == record
    before = list(load_workbook(path)["S30"].values)
    syncer.sync(path)
    assert list(load_workbook(path)["S30"].values) == before


def test_conflicting_filled_budgets_are_not_silently_overwritten(isolated_tables):
    path = isolated_tables / "Supplementary Tables.xlsx"
    csv_path = isolated_tables / "S46_training_run_budgets.csv"
    fields, rows = syncer.read_table(csv_path)
    next(r for r in rows if r["run"] == "hEvo-open-base")["optimizer_step_budget"] = 321
    syncer.write_table(csv_path, fields, rows)
    wb = load_workbook(path)
    wb["S30"].cell(public_row(wb["S30"]), 5, 999)
    wb.save(path)
    original = path.read_bytes()
    with pytest.raises(ValueError, match="Conflicting"):
        syncer.sync(path)
    assert path.read_bytes() == original
